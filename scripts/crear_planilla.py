#!/usr/bin/env python3
"""Genera el control financiero de la automotora (Excel / Google Sheets)."""

import sys
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from bot.catalogo import AFECTO_POR_DEFECTO, GASTOS_FIJOS, GASTOS_VARIABLES
from bot.formulas import suma_iva_concepto, suma_iva_tipo, suma_neto
from bot.planilla import Movimiento, preparar_hoja_comprobantes, _escribir_comprobante

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import SeriesLabel
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Protection,
    Side,
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.page import PageMargins
from openpyxl.comments import Comment
from openpyxl.workbook.defined_name import DefinedName

OUT = "Control_Financiero_Automotora.xlsx"

# Paleta
NAVY = "1A365D"
NAVY_DEEP = "0F2744"
GOLD = "C9A227"
CREAM = "FAF7F2"
INPUT_BG = "FFF8DC"
FORMULA_BG = "EDF2F7"
SECTION = "2C5282"
WHITE = "FFFFFF"
GREEN = "276749"
GREEN_BG = "C6F6D5"
RED = "C53030"
RED_BG = "FED7D7"
LINE = "CBD5E0"
ALT = "F7FAFC"
MUTED = "718096"
TEAL = "234E52"

MESES = [
    "Enero",
    "Febrero",
    "Marzo",
    "Abril",
    "Mayo",
    "Junio",
    "Julio",
    "Agosto",
    "Septiembre",
    "Octubre",
    "Noviembre",
    "Diciembre",
]

# Celdas reales de Configuracion: B5 nombre, B7 año, B8 mes, B11 IVA, B12 PPM, B13 renta
IVA = "Configuracion!$B$11"
PPM = "Configuracion!$B$12"
RENTA = "Configuracion!$B$13"
ANIO = "Configuracion!$B$7"
MES = "Configuracion!$B$8"
NOMBRE = "Configuracion!$B$5"

VA = "Ventas_Autos"
TR = "Taller_Repuestos"
OI = "Otros_Ingresos"
GF = "Gastos_Fijos"
GV = "Gastos_Variables"
ER = "Estado_Resultados"
SH_IVA = "IVA"
SH_PPM = "PPM"

VA_FIRST, VA_LAST = 6, 205
TR_FIRST, TR_LAST = 6, 205
OI_FIRST, OI_LAST = 6, 105
GF_FIRST, GF_LAST = 6, 25
GV_FIRST, GV_LAST = 6, 22

# Estado_Resultados rows
ER_HDR = 4
ER_ING_NUEVOS = 7
ER_ING_USADOS = 8
ER_ING_CONS = 9
ER_ING_TALLER = 10
ER_ING_REP = 11
ER_ING_OTROS = 12
ER_ING_TOTAL = 13
ER_COS_AUTOS = 16
ER_COS_TALLER = 17
ER_COS_TOTAL = 18
ER_MARGEN = 19
ER_MARGEN_PCT = 20
ER_GAS_FIJOS = 23
ER_GAS_COMIS = 24
ER_GAS_VAR = 25
ER_GAS_TOTAL = 26
ER_UTIL_OP = 27
ER_UTIL_PCT = 28
ER_IMP_RENTA = 31
ER_UTIL_NETA = 32

thin = Border(
    left=Side(style="thin", color=LINE),
    right=Side(style="thin", color=LINE),
    top=Side(style="thin", color=LINE),
    bottom=Side(style="thin", color=LINE),
)
thick_gold = Border(
    left=Side(style="medium", color=GOLD),
    right=Side(style="medium", color=GOLD),
    top=Side(style="medium", color=GOLD),
    bottom=Side(style="medium", color=GOLD),
)
side_thin = Side(style="thin", color=LINE)

font_title = Font(name="Calibri", size=20, bold=True, color=WHITE)
font_h = Font(name="Calibri", size=11, bold=True, color=WHITE)
font_section = Font(name="Calibri", size=11, bold=True, color=WHITE)
font_label = Font(name="Calibri", size=10, bold=True, color=NAVY)
font_normal = Font(name="Calibri", size=10, color="1A202C")
font_muted = Font(name="Calibri", size=9, italic=True, color=MUTED)
font_kpi = Font(name="Calibri", size=16, bold=True, color=NAVY_DEEP)
font_input = Font(name="Calibri", size=10, color="744210")
font_formula = Font(name="Calibri", size=10, color="2A4365")

fill_navy = PatternFill("solid", fgColor=NAVY)
fill_deep = PatternFill("solid", fgColor=NAVY_DEEP)
fill_gold = PatternFill("solid", fgColor=GOLD)
fill_section = PatternFill("solid", fgColor=SECTION)
fill_cream = PatternFill("solid", fgColor=CREAM)
fill_input = PatternFill("solid", fgColor=INPUT_BG)
fill_formula = PatternFill("solid", fgColor=FORMULA_BG)
fill_green = PatternFill("solid", fgColor=GREEN_BG)
fill_red = PatternFill("solid", fgColor=RED_BG)
fill_alt = PatternFill("solid", fgColor=ALT)
fill_white = PatternFill("solid", fgColor=WHITE)
fill_teal = PatternFill("solid", fgColor=TEAL)

center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left = Alignment(horizontal="left", vertical="center", wrap_text=True)
right = Alignment(horizontal="right", vertical="center")

FMT_CLP = '"$"#,##0'
FMT_PCT = "0.00%"
FMT_N = "0.00%"


def money(cell):
    cell.number_format = FMT_CLP


def pct(cell):
    cell.number_format = FMT_PCT


def apply_border_range(ws, r1, c1, r2, c2):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            ws.cell(r, c).border = thin


def header_bar(ws, row, cols, title, fill=fill_navy, size=18):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=cols)
    cell = ws.cell(row, 1, title)
    cell.font = Font(name="Calibri", size=size, bold=True, color=WHITE)
    cell.fill = fill
    cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for c in range(1, cols + 1):
        ws.cell(row, c).fill = fill
        ws.cell(row, c).font = Font(name="Calibri", size=size, bold=True, color=WHITE)
    ws.row_dimensions[row].height = 28


def sub_bar(ws, row, cols, text):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=cols)
    ws.cell(row, 1, text).font = Font(name="Calibri", size=10, color=WHITE)
    for c in range(1, cols + 1):
        ws.cell(row, c).fill = fill_deep
        ws.cell(row, c).font = Font(name="Calibri", size=10, color="E2E8F0")
        ws.cell(row, c).alignment = left
    ws.row_dimensions[row].height = 20


def section_row(ws, row, cols, text, fill=fill_section):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=cols)
    for c in range(1, cols + 1):
        ws.cell(row, c).fill = fill
        ws.cell(row, c).font = font_section
        ws.cell(row, c).alignment = left
    ws.cell(row, 1, text)
    ws.row_dimensions[row].height = 20


def style_header_row(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row, c)
        cell.font = font_h
        cell.fill = fill_navy
        cell.alignment = center
        cell.border = thin
    ws.row_dimensions[row].height = 32
    ws.auto_filter.ref = None  # set later


def input_cell(cell, value=None):
    if value is not None:
        cell.value = value
    cell.fill = fill_input
    cell.font = font_input
    cell.border = thin
    cell.alignment = center


def formula_cell(cell, formula, fmt=None):
    cell.value = formula
    cell.fill = fill_formula
    cell.font = font_formula
    cell.border = thin
    cell.alignment = Alignment(horizontal="center", vertical="center")
    if fmt == "clp":
        money(cell)
    elif fmt == "pct":
        pct(cell)


def page_setup(ws, landscape=True, fit=True):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    if fit:
        ws.page_setup.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5)
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 100


def month_col(m):
    """1 = Enero -> column B on ER/IVA/PPM (col 2)."""
    return m + 1


def add_list_dv(ws, formula, cells, prompt="Elija un valor"):
    dv = DataValidation(
        type="list",
        formula1=formula,
        allow_blank=True,
        showDropDown=False,
        showErrorMessage=True,
        errorTitle="Valor no válido",
        error="Use un valor de la lista.",
        promptTitle="Lista",
        prompt=prompt,
        showInputMessage=True,
    )
    ws.add_data_validation(dv)
    dv.add(cells)


def build_instrucciones(ws):
    page_setup(ws, landscape=False, fit=False)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 92
    header_bar(ws, 1, 2, "  Cómo usar esta planilla", size=18)
    sub_bar(ws, 2, 2, "  Control de ingresos, gastos, IVA y PPM — Automotora")

    bloques = [
        (
            4,
            "Color de las celdas",
            "Amarillo claro: tú escribes. Gris azulado: fórmula, no la borres. "
            "Los totales y el resumen se calculan solos.",
        ),
        (
            7,
            "Orden de trabajo cada mes",
            "1) En Configuracion elige el año y el mes que estás cerrando.\n"
            "2) Carga ventas de autos (propios y consignación) en Ventas_Autos.\n"
            "3) Carga taller y repuestos en Taller_Repuestos.\n"
            "4) Los gastos entran por la hoja Comprobantes (bot de Telegram) y se suman solos en el mes.\n"
            "5) Revisa Resumen, IVA y PPM. El F29 del mes se declara al mes siguiente.",
        ),
        (
            10,
            "IVA (19%)",
            "El IVA no es un gasto de la automotora: es un impuesto que cobras y pagas.\n"
            "• IVA débito: el IVA de tus ventas y servicios (autos vendidos, taller, repuestos, comisiones).\n"
            "• IVA crédito: el IVA de tus compras afectas (autos con factura, insumos, publicidad, luz, etc.).\n"
            "• IVA a pagar = débito − crédito − remanente del mes anterior. Si da negativo, queda remanente.\n"
            "Importante: si compras un usado a un particular, casi nunca hay crédito fiscal. Marca Compra afecta IVA = No.",
        ),
        (
            13,
            "PPM (Pagos Provisionales Mensuales)",
            "El PPM se calcula sobre los ingresos netos del mes (sin IVA), no sobre la utilidad. "
            "Aunque un mes salgas a pérdida, igual puede corresponder PPM.\n"
            "La tasa vive en Configuracion. Si es el primer año comercial suele usarse 0,25%. "
            "Después el SII fija un factor según el año anterior. Confírmala con tu contador.\n"
            "El PPM se imputa al impuesto a la renta de primera categoría al cierre del año.",
        ),
        (
            16,
            "Consignación vs. venta propia",
            "• Auto propio (Nuevo o Usado): el ingreso es el precio de venta. El costo es lo que te costó. "
            "La utilidad de la unidad es venta neta − costo neto − comisión del vendedor.\n"
            "• Consignación: el auto no es tuyo. El ingreso de la automotora es la comisión. "
            "Llena Comisión consignación (con IVA) y deja el costo en 0.",
        ),
        (
            19,
            "Pasar a Google Sheets",
            "1. Sube este archivo a Google Drive.\n"
            "2. Clic derecho → Abrir con → Google Sheets.\n"
            "3. Archivo → Guardar como Google Sheets (así queda nativo y se puede compartir).\n"
            "Las fórmulas están en formato compatible (SUMIFS, INDEX, IF). No uses macros.",
        ),
        (
            22,
            "Qué no reemplaza esta planilla",
            "Es un control de gestión para ver si el negocio gana plata y qué hay que pagar de IVA y PPM. "
            "No reemplaza el libro de compras/ventas ni el F29 que hace tu contador. "
            "Si hay diferencias con el SII, manda la planilla y los documentos al contador.",
        ),
    ]
    r = 4
    for title, body in [
        (b[1], b[2]) for b in bloques
    ]:
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=2)
        ws.cell(r, 2, title).font = Font(name="Calibri", size=13, bold=True, color=NAVY)
        ws.cell(r, 2).fill = fill_cream
        ws.row_dimensions[r].height = 22
        r += 1
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=2)
        ws.cell(r, 2, body).font = font_normal
        ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 78 if len(body) > 180 else 64
        r += 2

    ws.cell(r, 2, "Las filas de ejemplo en agosto 2026 se pueden borrar. No borres columnas ni encabezados.").font = font_muted
    ws.freeze_panes = "A3"
    ws.sheet_properties.tabColor = GOLD


def build_config(ws):
    page_setup(ws, landscape=False)
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 62
    header_bar(ws, 1, 3, "  Configuración", size=18)
    sub_bar(ws, 2, 3, "  Cambia solo las celdas amarillas. El resto de la planilla toma estos valores.")

    ws.cell(4, 1, "DATOS DE LA EMPRESA").font = font_label
    labels = [
        (5, "Nombre de la automotora", "Mi Automotora", "Aparece en el resumen y en las impresiones."),
        (6, "RUT", "76.000.000-K", "Solo referencia; no se usa en fórmulas."),
        (7, "Año de trabajo", 2026, "Filtra ventas, IVA y PPM. Cámbialo en enero."),
        (8, "Mes de trabajo (1 a 12)", 8, "El Resumen, IVA y PPM muestran este mes. 8 = agosto."),
    ]
    for row, lab, val, hint in labels:
        ws.cell(row, 1, lab).font = font_label
        ws.cell(row, 1).alignment = left
        input_cell(ws.cell(row, 2), val)
        ws.cell(row, 3, hint).font = font_muted
        ws.cell(row, 3).alignment = left

    ws.cell(7, 2).number_format = "0"
    ws.cell(8, 2).number_format = "0"
    add_list_dv(ws, '"1,2,3,4,5,6,7,8,9,10,11,12"', "B8")

    ws.cell(10, 1, "TASAS TRIBUTARIAS").font = font_label
    tax = [
        (11, "Tasa IVA", 0.19, "Chile: 19%. No la cambies salvo cambio legal."),
        (12, "Tasa PPM", 0.0025, "Por defecto 0,25% (primer año). Reemplázala por el factor del SII."),
        (13, "Tasa impuesto 1ª categoría", 0.27, "27% régimen general. PyME 14 D puede ser 25%. Confirma con tu contador."),
    ]
    for row, lab, val, hint in tax:
        ws.cell(row, 1, lab).font = font_label
        input_cell(ws.cell(row, 2), val)
        pct(ws.cell(row, 2))
        ws.cell(row, 3, hint).font = font_muted

    ws.cell(15, 1, "MES EN CURSO (se calcula solo)").font = font_label
    formula_cell(ws.cell(16, 2), f'=INDEX({{"Enero","Febrero","Marzo","Abril","Mayo","Junio","Julio","Agosto","Septiembre","Octubre","Noviembre","Diciembre"}},B8)')
    ws.cell(16, 1, "Nombre del mes").font = font_label
    ws.cell(16, 3, "Lo usa el Resumen como título del período.").font = font_muted
    ws.cell(16, 2).alignment = center

    # Array constant INDEX can be flaky in Sheets. Safer: CHOOSE
    ws["B16"].value = '=CHOOSE(B8,"Enero","Febrero","Marzo","Abril","Mayo","Junio","Julio","Agosto","Septiembre","Octubre","Noviembre","Diciembre")'

    ws.cell(18, 1, "LEYENDA").font = font_label
    input_cell(ws.cell(19, 2), "Dato que tú escribes")
    ws.cell(19, 1, "Celda de entrada").font = font_normal
    formula_cell(ws.cell(20, 2), "=\"Cálculo automático\"")
    ws.cell(20, 1, "Celda calculada").font = font_normal

    ws.cell(22, 1, "Notas para el contador").font = font_label
    ws.merge_cells("B22:C24")
    input_cell(ws.cell(22, 2), "Tasa PPM confirmada con contador: _  /  Régimen: _")
    ws.cell(22, 2).alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
    ws.merge_cells("A22:A24")
    ws.row_dimensions[22].height = 36

    ws.merge_cells("A26:C28")
    ws.cell(26, 1).value = (
        "Los movimientos de agosto 2026 son EJEMPLOS (pocas ventas y una nómina completa, "
        "por eso el mes puede salir a pérdida). Bórralos y carga tus datos reales. "
        "Cambia el nombre de la automotora arriba."
    )
    ws.cell(26, 1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(26, 1).font = font_muted
    ws.row_dimensions[26].height = 48

    ws.freeze_panes = "A4"
    ws.sheet_properties.tabColor = GOLD

    ws["B5"].comment = Comment("Nombre comercial o razón social.", "Sistema")
    ws["B12"].comment = Comment(
        "PPM = ingresos netos del mes × esta tasa. No se calcula sobre la utilidad.",
        "Sistema",
    )


def fill_month_headers(ws, row, start_col=2):
    ws.cell(row, 1, "Concepto").font = font_h
    ws.cell(row, 1).fill = fill_navy
    ws.cell(row, 1).alignment = center
    ws.cell(row, 1).border = thin
    for i, m in enumerate(MESES):
        cell = ws.cell(row, start_col + i, m)
        cell.font = font_h
        cell.fill = fill_navy
        cell.alignment = center
        cell.border = thin
    total = ws.cell(row, start_col + 12, "Total año")
    total.font = font_h
    total.fill = fill_gold
    total.alignment = center
    total.font = Font(name="Calibri", size=11, bold=True, color=NAVY_DEEP)
    total.border = thin


def build_ventas(ws):
    page_setup(ws, landscape=True)
    ncols = 26
    header_bar(ws, 1, ncols, "  Ventas de autos — nuevos, usados y consignación")
    sub_bar(
        ws,
        2,
        ncols,
        "  Amarillo = datos. En consignación el ingreso es la comisión. En propios, el precio de venta. "
        "Estado Vendido entra al resultado; Stock no.",
    )
    ws.merge_cells("A3:Z3")
    ws.cell(
        3,
        1,
        "IVA débito sale de la venta (o de la comisión). IVA crédito sale de la compra solo si Compra afecta IVA = Sí.",
    ).font = font_muted

    headers = [
        "N°",
        "Fecha compra",
        "Fecha venta",
        "Tipo",
        "Marca",
        "Modelo",
        "Año auto",
        "Patente",
        "Proveedor / consignante",
        "Compra afecta IVA",
        "Costo neto",
        "IVA compra",
        "Total compra",
        "Precio venta (c/IVA)",
        "Comisión consignación (c/IVA)",
        "Comisión vendedor (neto)",
        "Estado",
        "Notas",
        "Año compra",
        "Mes compra",
        "Año venta",
        "Mes venta",
        "Ingreso neto automotora",
        "IVA débito",
        "Costo en resultado",
        "Margen de la unidad",
    ]
    for i, h in enumerate(headers, 1):
        cell = ws.cell(5, i, h)
        cell.font = font_h
        cell.fill = fill_navy if i < 19 else fill_teal
        cell.alignment = center
        cell.border = thin
    ws.row_dimensions[5].height = 36

    samples = [
        {
            "b": date(2026, 7, 18),
            "c": date(2026, 8, 5),
            "d": "Usado",
            "e": "Toyota",
            "f": "Corolla",
            "g": 2019,
            "h": "RPKB12",
            "i": "Particular",
            "j": "No",
            "k": 8500000,
            "n": 11900000,
            "o": None,
            "p": 280000,
            "q": "Vendido",
            "r": "EJEMPLO — usado comprado a particular (sin crédito IVA)",
        },
        {
            "b": date(2026, 8, 2),
            "c": date(2026, 8, 20),
            "d": "Nuevo",
            "e": "Chevrolet",
            "f": "Onix",
            "g": 2026,
            "h": "JKLP34",
            "i": "Importadora / marca",
            "j": "Sí",
            "k": 10500000,
            "n": 14900000,
            "o": None,
            "p": 350000,
            "q": "Vendido",
            "r": "EJEMPLO — nuevo con factura (sí hay crédito IVA)",
        },
        {
            "b": date(2026, 8, 1),
            "c": date(2026, 8, 12),
            "d": "Consignación",
            "e": "Hyundai",
            "f": "Tucson",
            "g": 2021,
            "h": "LFTW90",
            "i": "Juan Pérez (consignante)",
            "j": "No",
            "k": 0,
            "n": 14500000,
            "o": 950000,
            "p": 120000,
            "q": "Vendido",
            "r": "EJEMPLO — ingreso = comisión, no el precio del auto",
        },
        {
            "b": date(2026, 8, 22),
            "c": None,
            "d": "Usado",
            "e": "Mazda",
            "f": "3",
            "g": 2018,
            "h": "PDKS21",
            "i": "Particular",
            "j": "No",
            "k": 7200000,
            "n": None,
            "o": None,
            "p": None,
            "q": "Stock",
            "r": "EJEMPLO — en patio, no entra al resultado hasta venderse",
        },
    ]

    for r in range(VA_FIRST, VA_LAST + 1):
        alt = fill_alt if r % 2 == 0 else fill_white
        for c in range(1, 27):
            ws.cell(r, c).border = thin
            ws.cell(r, c).font = font_normal
            ws.cell(r, c).alignment = center

        formula_cell(ws.cell(r, 1), f'=IF(D{r}="","",ROW()-5)')
        for c in (2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 14, 15, 16, 17, 18):
            ws.cell(r, c).fill = fill_input
            ws.cell(r, c).font = font_input
        ws.cell(r, 7).number_format = "0"
        money(ws.cell(r, 11))
        money(ws.cell(r, 14))
        money(ws.cell(r, 15))
        money(ws.cell(r, 16))
        ws.cell(r, 2).number_format = "DD/MM/YYYY"
        ws.cell(r, 3).number_format = "DD/MM/YYYY"

        formula_cell(ws.cell(r, 12), f'=IF(K{r}="","",IF(J{r}="Sí",ROUND(K{r}*{IVA},0),0))', "clp")
        formula_cell(ws.cell(r, 13), f'=IF(K{r}="","",K{r}+L{r})', "clp")
        formula_cell(ws.cell(r, 19), f'=IF(B{r}="","",YEAR(B{r}))')
        formula_cell(ws.cell(r, 20), f'=IF(B{r}="","",MONTH(B{r}))')
        formula_cell(ws.cell(r, 21), f'=IF(C{r}="","",YEAR(C{r}))')
        formula_cell(ws.cell(r, 22), f'=IF(C{r}="","",MONTH(C{r}))')
        formula_cell(
            ws.cell(r, 23),
            f'=IF(D{r}="","",IF(D{r}="Consignación",IF(O{r}="",0,ROUND(O{r}/(1+{IVA}),0)),IF(N{r}="",0,ROUND(N{r}/(1+{IVA}),0))))',
            "clp",
        )
        formula_cell(ws.cell(r, 24), f'=IF(W{r}="","",ROUND(W{r}*{IVA},0))', "clp")
        formula_cell(
            ws.cell(r, 25),
            f'=IF(OR(D{r}="",Q{r}<>"Vendido"),0,IF(D{r}="Consignación",0,IF(K{r}="",0,K{r})))',
            "clp",
        )
        formula_cell(
            ws.cell(r, 26),
            f'=IF(Q{r}<>"Vendido","",W{r}-Y{r}-IF(P{r}="",0,P{r}))',
            "clp",
        )

        # default fills for unused sample rows stay input; override formula cols already set
        if r >= VA_FIRST + len(samples):
            for c in range(2, 19):
                if c not in (12, 13):
                    ws.cell(r, c).fill = fill_input

    for i, s in enumerate(samples):
        r = VA_FIRST + i
        ws.cell(r, 2).value = s["b"]
        if s["c"]:
            ws.cell(r, 3).value = s["c"]
        ws.cell(r, 4).value = s["d"]
        ws.cell(r, 5).value = s["e"]
        ws.cell(r, 6).value = s["f"]
        ws.cell(r, 7).value = s["g"]
        ws.cell(r, 8).value = s["h"]
        ws.cell(r, 9).value = s["i"]
        ws.cell(r, 10).value = s["j"]
        ws.cell(r, 11).value = s["k"]
        if s["n"] is not None:
            ws.cell(r, 14).value = s["n"]
        if s["o"] is not None:
            ws.cell(r, 15).value = s["o"]
        if s["p"] is not None:
            ws.cell(r, 16).value = s["p"]
        ws.cell(r, 17).value = s["q"]
        ws.cell(r, 18).value = s["r"]

    widths = {
        "A": 6, "B": 14, "C": 14, "D": 14, "E": 14, "F": 14, "G": 10,
        "H": 12, "I": 24, "J": 16, "K": 14, "L": 13, "M": 14, "N": 18,
        "O": 22, "P": 20, "Q": 16, "R": 42, "S": 12, "T": 12, "U": 12,
        "V": 12, "W": 22, "X": 13, "Y": 16, "Z": 18,
    }
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    add_list_dv(ws, '"Nuevo,Usado,Consignación"', f"D{VA_FIRST}:D{VA_LAST}")
    add_list_dv(ws, '"Sí,No"', f"J{VA_FIRST}:J{VA_LAST}")
    add_list_dv(ws, '"Stock,Vendido,Reservado,En consignación,Devuelto"', f"Q{VA_FIRST}:Q{VA_LAST}")

    ws.auto_filter.ref = f"A5:Z{VA_LAST}"
    ws.freeze_panes = "E6"
    ws.sheet_properties.tabColor = "38A169"

    ws.conditional_formatting.add(
        f"Z{VA_FIRST}:Z{VA_LAST}",
        CellIsRule(operator="lessThan", formula=["0"], fill=fill_red, font=Font(color=RED, bold=True)),
    )
    ws.conditional_formatting.add(
        f"Z{VA_FIRST}:Z{VA_LAST}",
        CellIsRule(operator="greaterThan", formula=["0"], fill=fill_green, font=Font(color=GREEN, bold=True)),
    )


def build_taller(ws):
    page_setup(ws)
    header_bar(ws, 1, 14, "  Taller y repuestos")
    sub_bar(ws, 2, 14, "  Mano de obra, OT y mostrador. Neto = sin IVA. El margen es venta neta − costo neto.")
    ws.merge_cells("A3:N3")
    ws.cell(3, 1, "Si el costo viene con factura, marca Costo afecto IVA = Sí para que sume al crédito fiscal.").font = font_muted

    headers = [
        "Fecha", "Tipo", "Documento / OT", "Cliente", "Descripción",
        "Venta neta", "IVA venta", "Total venta", "Costo neto",
        "Costo afecto IVA", "IVA crédito", "Margen", "Año", "Mes",
    ]
    for i, h in enumerate(headers, 1):
        cell = ws.cell(5, i, h)
        cell.font = font_h
        cell.fill = fill_navy if i <= 12 else fill_teal
        cell.alignment = center
        cell.border = thin
    ws.row_dimensions[5].height = 32

    samples = [
        (date(2026, 8, 8), "Taller", "OT-0145", "Cliente taller", "Service 10.000 km", 350000, 80000, "Sí"),
        (date(2026, 8, 15), "Repuestos", "FAC-889", "Mostrador", "Filtro + pastillas", 180000, 95000, "Sí"),
        (date(2026, 8, 21), "Taller", "OT-0152", "Preparación unidad", "Pintura paragolpes", 220000, 110000, "Sí"),
    ]
    for r in range(TR_FIRST, TR_LAST + 1):
        for c in range(1, 15):
            ws.cell(r, c).border = thin
            ws.cell(r, c).alignment = center
        for c in (1, 2, 3, 4, 5, 6, 9, 10):
            ws.cell(r, c).fill = fill_input
            ws.cell(r, c).font = font_input
        ws.cell(r, 1).number_format = "DD/MM/YYYY"
        money(ws.cell(r, 6))
        money(ws.cell(r, 9))
        formula_cell(ws.cell(r, 7), f'=IF(F{r}="","",ROUND(F{r}*{IVA},0))', "clp")
        formula_cell(ws.cell(r, 8), f'=IF(F{r}="","",F{r}+G{r})', "clp")
        formula_cell(ws.cell(r, 11), f'=IF(I{r}="","",IF(J{r}="Sí",ROUND(I{r}*{IVA},0),0))', "clp")
        formula_cell(ws.cell(r, 12), f'=IF(F{r}="","",F{r}-IF(I{r}="",0,I{r}))', "clp")
        formula_cell(ws.cell(r, 13), f'=IF(A{r}="","",YEAR(A{r}))')
        formula_cell(ws.cell(r, 14), f'=IF(A{r}="","",MONTH(A{r}))')

    for i, s in enumerate(samples):
        r = TR_FIRST + i
        ws.cell(r, 1).value = s[0]
        ws.cell(r, 2).value = s[1]
        ws.cell(r, 3).value = s[2]
        ws.cell(r, 4).value = s[3]
        ws.cell(r, 5).value = s[4]
        ws.cell(r, 6).value = s[5]
        ws.cell(r, 9).value = s[6]
        ws.cell(r, 10).value = s[7]

    widths = {"A": 13, "B": 12, "C": 16, "D": 22, "E": 28, "F": 14, "G": 12, "H": 14, "I": 13, "J": 16, "K": 13, "L": 13, "M": 10, "N": 8}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    add_list_dv(ws, '"Taller,Repuestos"', f"B{TR_FIRST}:B{TR_LAST}")
    add_list_dv(ws, '"Sí,No"', f"J{TR_FIRST}:J{TR_LAST}")
    ws.auto_filter.ref = f"A5:N{TR_LAST}"
    ws.freeze_panes = "A6"
    ws.sheet_properties.tabColor = "38A169"

    ws.conditional_formatting.add(
        f"L{TR_FIRST}:L{TR_LAST}",
        CellIsRule(operator="lessThan", formula=["0"], fill=fill_red),
    )


def build_otros(ws):
    page_setup(ws)
    header_bar(ws, 1, 9, "  Otros ingresos")
    sub_bar(ws, 2, 9, "  Financiamiento, transferencias cobradas al cliente, arriendo de patio, etc.")
    headers = ["Fecha", "Categoría", "Descripción", "Neto", "Afecto IVA", "IVA", "Total", "Año", "Mes"]
    for i, h in enumerate(headers, 1):
        cell = ws.cell(5, i, h)
        cell.font = font_h
        cell.fill = fill_navy if i <= 7 else fill_teal
        cell.alignment = center
        cell.border = thin
    for r in range(OI_FIRST, OI_LAST + 1):
        for c in range(1, 10):
            ws.cell(r, c).border = thin
            ws.cell(r, c).alignment = center
        for c in (1, 2, 3, 4, 5):
            ws.cell(r, c).fill = fill_input
            ws.cell(r, c).font = font_input
        ws.cell(r, 1).number_format = "DD/MM/YYYY"
        money(ws.cell(r, 4))
        formula_cell(ws.cell(r, 6), f'=IF(D{r}="","",IF(E{r}="Sí",ROUND(D{r}*{IVA},0),0))', "clp")
        formula_cell(ws.cell(r, 7), f'=IF(D{r}="","",D{r}+F{r})', "clp")
        formula_cell(ws.cell(r, 8), f'=IF(A{r}="","",YEAR(A{r}))')
        formula_cell(ws.cell(r, 9), f'=IF(A{r}="","",MONTH(A{r}))')

    ws.cell(6, 1).value = date(2026, 8, 18)
    ws.cell(6, 2).value = "Gestión de financiamiento"
    ws.cell(6, 3).value = "EJEMPLO — comisión financiera"
    ws.cell(6, 4).value = 120000
    ws.cell(6, 5).value = "Sí"

    for col, w in zip("ABCDEFGHI", [13, 26, 36, 14, 14, 12, 14, 10, 8]):
        ws.column_dimensions[col].width = w
    add_list_dv(
        ws,
        '"Gestión de financiamiento,Transferencia,Arriendo espacio,Otros"',
        f"B{OI_FIRST}:B{OI_LAST}",
    )
    add_list_dv(ws, '"Sí,No"', f"E{OI_FIRST}:E{OI_LAST}")
    ws.auto_filter.ref = f"A5:I{OI_LAST}"
    ws.freeze_panes = "A6"
    ws.sheet_properties.tabColor = "38A169"


def build_gastos(ws, titulo, sub, items, tab="C53030"):
    page_setup(ws)
    header_bar(ws, 1, 18, f"  {titulo}")
    sub_bar(ws, 2, 18, f"  {sub}")
    ws.merge_cells("A3:R3")
    tipo = "Gasto fijo" if titulo.startswith("Gastos fijos") else "Gasto variable"
    ws.cell(
        3,
        1,
        "Los montos de cada mes salen de Comprobantes (el bot de Telegram). No escribas en las columnas de los meses. "
        "Afecto IVA de la columna B es la referencia del concepto; el crédito real sale de cada comprobante.",
    ).font = font_muted

    headers = ["Concepto", "Afecto IVA"] + MESES + ["Total año", "IVA del año", "Este mes", "IVA este mes"]
    for i, h in enumerate(headers, 1):
        cell = ws.cell(5, i, h)
        cell.font = Font(name="Calibri", size=10, bold=True, color=WHITE if i <= 14 else NAVY_DEEP)
        cell.fill = fill_navy if i <= 14 else fill_gold
        cell.alignment = center
        cell.border = thin
    ws.row_dimensions[5].height = 32

    for idx, (concepto, afecto) in enumerate(items):
        r = GF_FIRST + idx if titulo.startswith("Gastos fijos") else GV_FIRST + idx
        # caller uses same first row 6
        r = 6 + idx
        ws.cell(r, 1, concepto).font = font_normal
        ws.cell(r, 1).alignment = left
        ws.cell(r, 1).fill = fill_input
        ws.cell(r, 1).border = thin
        input_cell(ws.cell(r, 2), afecto)
        for m in range(12):
            cell = ws.cell(r, 3 + m)
            formula_cell(cell, suma_neto(tipo, r, m + 1), "clp")
        formula_cell(ws.cell(r, 15), f"=SUM(C{r}:N{r})", "clp")
        formula_cell(ws.cell(r, 16), suma_iva_concepto(tipo, r), "clp")
        formula_cell(ws.cell(r, 17), f"=IFERROR(INDEX(C{r}:N{r},1,{MES}),0)", "clp")
        formula_cell(ws.cell(r, 18), suma_iva_concepto(tipo, r, "actual"), "clp")

    last = 5 + len(items)
    # total row
    tot = last + 1
    ws.cell(tot, 1, "TOTAL").font = Font(name="Calibri", size=11, bold=True, color=WHITE)
    ws.cell(tot, 1).fill = fill_deep
    ws.cell(tot, 2).fill = fill_deep
    ws.cell(tot, 2).border = thin
    for c in range(1, 19):
        ws.cell(tot, c).fill = fill_deep
        ws.cell(tot, c).font = Font(name="Calibri", size=10, bold=True, color=WHITE)
        ws.cell(tot, c).border = thin
    for c in range(3, 19):
        col = get_column_letter(c)
        ws.cell(tot, c).value = f"=SUM({col}6:{col}{last})"
        money(ws.cell(tot, c))
        ws.cell(tot, c).fill = fill_deep
        ws.cell(tot, c).font = Font(name="Calibri", size=10, bold=True, color=WHITE)

    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 12
    for i in range(3, 19):
        ws.column_dimensions[get_column_letter(i)].width = 13
    add_list_dv(ws, '"Sí,No"', f"B6:B{last}")
    ws.freeze_panes = "C6"
    ws.sheet_properties.tabColor = tab
    ws.auto_filter.ref = f"A5:R{last}"
    return tot


def build_er(ws):
    page_setup(ws)
    header_bar(ws, 1, 14, "  Estado de resultados — 12 meses")
    sub_bar(ws, 2, 14, "  Motor de cálculo. No edites esta hoja: se alimenta de las cargas. Cifras NETAS, sin IVA.")
    ws.cell(3, 1, "Año de las cifras").font = font_muted
    formula_cell(ws.cell(3, 2), f"={ANIO}")
    ws.cell(3, 2).fill = fill_cream
    ws.merge_cells("C3:N3")
    ws.cell(3, 3, "Todas las columnas de esta hoja son fórmulas. Los datos se cargan en las otras pestañas.").font = font_muted

    fill_month_headers(ws, ER_HDR)
    ws.row_dimensions[ER_HDR].height = 22

    def msumifs(sheet, value_col, year_col, month_col, month_num, extra=""):
        return (
            f"SUMIFS({sheet}!${value_col}${start}:${value_col}${end},"
            f"{sheet}!${year_col}${start}:${year_col}${end},{ANIO},"
            f"{sheet}!${month_col}${start}:${month_col}${end},{month_num}"
            f"{extra})"
        )

    # helpers per month column
    def set_row(row, label, kind, fill=None, bold=False):
        cell = ws.cell(row, 1, label)
        cell.alignment = left
        cell.border = thin
        cell.font = Font(name="Calibri", size=10, bold=bold, color=NAVY if bold else "1A202C")
        if fill:
            cell.fill = fill
        for m in range(1, 13):
            c = month_col(m)
            ws.cell(row, c).border = thin
            if fill:
                ws.cell(row, c).fill = fill
            money(ws.cell(row, c))
        tot = ws.cell(row, 14)
        tot.value = f"=SUM(B{row}:M{row})"
        tot.border = thin
        tot.fill = fill_gold if not fill else fill
        tot.font = Font(name="Calibri", size=10, bold=True, color=NAVY_DEEP)
        money(tot)

    def write_sumifs_row(row, label, formulas_by_month, fill=None):
        set_row(row, label, "f", fill)
        for m in range(1, 13):
            cell = ws.cell(row, month_col(m))
            cell.value = formulas_by_month[m]
            cell.fill = fill_formula if not fill else fill
            cell.font = font_formula
            money(cell)

    section_row(ws, 6, 14, "INGRESOS")
    # Autos nuevos
    f_nuevos = {}
    f_usados = {}
    f_cons = {}
    f_taller = {}
    f_rep = {}
    f_otros = {}
    f_cos_autos = {}
    f_cos_taller = {}
    f_comis = {}
    for m in range(1, 13):
        f_nuevos[m] = (
            f'=SUMIFS({VA}!$W${VA_FIRST}:$W${VA_LAST},{VA}!$D${VA_FIRST}:$D${VA_LAST},"Nuevo",'
            f'{VA}!$U${VA_FIRST}:$U${VA_LAST},{ANIO},{VA}!$V${VA_FIRST}:$V${VA_LAST},{m},'
            f'{VA}!$Q${VA_FIRST}:$Q${VA_LAST},"Vendido")'
        )
        f_usados[m] = (
            f'=SUMIFS({VA}!$W${VA_FIRST}:$W${VA_LAST},{VA}!$D${VA_FIRST}:$D${VA_LAST},"Usado",'
            f'{VA}!$U${VA_FIRST}:$U${VA_LAST},{ANIO},{VA}!$V${VA_FIRST}:$V${VA_LAST},{m},'
            f'{VA}!$Q${VA_FIRST}:$Q${VA_LAST},"Vendido")'
        )
        f_cons[m] = (
            f'=SUMIFS({VA}!$W${VA_FIRST}:$W${VA_LAST},{VA}!$D${VA_FIRST}:$D${VA_LAST},"Consignación",'
            f'{VA}!$U${VA_FIRST}:$U${VA_LAST},{ANIO},{VA}!$V${VA_FIRST}:$V${VA_LAST},{m},'
            f'{VA}!$Q${VA_FIRST}:$Q${VA_LAST},"Vendido")'
        )
        f_taller[m] = (
            f'=SUMIFS({TR}!$F${TR_FIRST}:$F${TR_LAST},{TR}!$B${TR_FIRST}:$B${TR_LAST},"Taller",'
            f'{TR}!$M${TR_FIRST}:$M${TR_LAST},{ANIO},{TR}!$N${TR_FIRST}:$N${TR_LAST},{m})'
        )
        f_rep[m] = (
            f'=SUMIFS({TR}!$F${TR_FIRST}:$F${TR_LAST},{TR}!$B${TR_FIRST}:$B${TR_LAST},"Repuestos",'
            f'{TR}!$M${TR_FIRST}:$M${TR_LAST},{ANIO},{TR}!$N${TR_FIRST}:$N${TR_LAST},{m})'
        )
        f_otros[m] = (
            f'=SUMIFS({OI}!$D${OI_FIRST}:$D${OI_LAST},{OI}!$H${OI_FIRST}:$H${OI_LAST},{ANIO},'
            f'{OI}!$I${OI_FIRST}:$I${OI_LAST},{m})'
        )
        f_cos_autos[m] = (
            f'=SUMIFS({VA}!$Y${VA_FIRST}:$Y${VA_LAST},{VA}!$U${VA_FIRST}:$U${VA_LAST},{ANIO},'
            f'{VA}!$V${VA_FIRST}:$V${VA_LAST},{m},{VA}!$Q${VA_FIRST}:$Q${VA_LAST},"Vendido")'
        )
        f_cos_taller[m] = (
            f'=SUMIFS({TR}!$I${TR_FIRST}:$I${TR_LAST},{TR}!$M${TR_FIRST}:$M${TR_LAST},{ANIO},'
            f'{TR}!$N${TR_FIRST}:$N${TR_LAST},{m})'
        )
        f_comis[m] = (
            f'=SUMIFS({VA}!$P${VA_FIRST}:$P${VA_LAST},{VA}!$U${VA_FIRST}:$U${VA_LAST},{ANIO},'
            f'{VA}!$V${VA_FIRST}:$V${VA_LAST},{m},{VA}!$Q${VA_FIRST}:$Q${VA_LAST},"Vendido")'
        )

    write_sumifs_row(ER_ING_NUEVOS, "Autos nuevos", f_nuevos)
    write_sumifs_row(ER_ING_USADOS, "Autos usados", f_usados)
    write_sumifs_row(ER_ING_CONS, "Comisiones por consignación", f_cons)
    write_sumifs_row(ER_ING_TALLER, "Taller (mano de obra)", f_taller)
    write_sumifs_row(ER_ING_REP, "Repuestos", f_rep)
    write_sumifs_row(ER_ING_OTROS, "Otros ingresos", f_otros)

    set_row(ER_ING_TOTAL, "TOTAL INGRESOS", "t", fill_navy, bold=True)
    ws.cell(ER_ING_TOTAL, 1).font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    for m in range(1, 13):
        c = month_col(m)
        col = get_column_letter(c)
        cell = ws.cell(ER_ING_TOTAL, c)
        cell.value = f"={col}{ER_ING_NUEVOS}+{col}{ER_ING_USADOS}+{col}{ER_ING_CONS}+{col}{ER_ING_TALLER}+{col}{ER_ING_REP}+{col}{ER_ING_OTROS}"
        cell.fill = fill_navy
        cell.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
        money(cell)
    ws.cell(ER_ING_TOTAL, 14).value = f"=SUM(B{ER_ING_TOTAL}:M{ER_ING_TOTAL})"
    ws.cell(ER_ING_TOTAL, 14).fill = fill_navy
    ws.cell(ER_ING_TOTAL, 14).font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    money(ws.cell(ER_ING_TOTAL, 14))

    section_row(ws, 15, 14, "COSTOS DIRECTOS")
    write_sumifs_row(ER_COS_AUTOS, "Costo de autos vendidos", f_cos_autos)
    write_sumifs_row(ER_COS_TALLER, "Costo taller y repuestos", f_cos_taller)
    set_row(ER_COS_TOTAL, "TOTAL COSTOS DIRECTOS", "t", fill=PatternFill("solid", fgColor="4A5568"), bold=True)
    ws.cell(ER_COS_TOTAL, 1).font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    for m in range(1, 13):
        c = month_col(m)
        col = get_column_letter(c)
        cell = ws.cell(ER_COS_TOTAL, c)
        cell.value = f"={col}{ER_COS_AUTOS}+{col}{ER_COS_TALLER}"
        cell.fill = PatternFill("solid", fgColor="4A5568")
        cell.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
        money(cell)
    ws.cell(ER_COS_TOTAL, 14).fill = PatternFill("solid", fgColor="4A5568")
    ws.cell(ER_COS_TOTAL, 14).font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    ws.cell(ER_COS_TOTAL, 14).value = f"=SUM(B{ER_COS_TOTAL}:M{ER_COS_TOTAL})"
    money(ws.cell(ER_COS_TOTAL, 14))

    set_row(ER_MARGEN, "MARGEN BRUTO", "t", fill_green, bold=True)
    for m in range(1, 13):
        c = month_col(m)
        col = get_column_letter(c)
        cell = ws.cell(ER_MARGEN, c)
        cell.value = f"={col}{ER_ING_TOTAL}-{col}{ER_COS_TOTAL}"
        cell.fill = fill_green
        cell.font = Font(name="Calibri", size=10, bold=True, color=GREEN)
        money(cell)
    ws.cell(ER_MARGEN, 14).value = f"=N{ER_ING_TOTAL}-N{ER_COS_TOTAL}"
    ws.cell(ER_MARGEN, 14).fill = fill_green
    money(ws.cell(ER_MARGEN, 14))
    ws.cell(ER_MARGEN, 1).fill = fill_green

    set_row(ER_MARGEN_PCT, "Margen bruto %", "p")
    for m in range(1, 14):
        col = get_column_letter(m + 1) if m < 13 else "N"
        src_c = m + 1 if m < 13 else 14
        cell = ws.cell(ER_MARGEN_PCT, src_c)
        cell.value = f'=IF({col}{ER_ING_TOTAL}=0,"",{col}{ER_MARGEN}/{col}{ER_ING_TOTAL})'
        pct(cell)
        cell.fill = fill_formula
        cell.border = thin
        cell.font = font_formula
    ws.cell(ER_MARGEN_PCT, 1).fill = fill_white

    section_row(ws, 22, 14, "GASTOS OPERACIONALES")
    # Gastos fijos: row 24 of GF is total if 18 items: 6+18-1=23, total=24. We'll compute dynamically after build.
    # We'll use named approach: Gastos_Fijos total row stored later. Use SUM of Este_mes column? 
    # For each month, SUM the month column C-N excluding header and including all item rows, not the total row to avoid double.
    # Items GF: 18 rows (6-23), total 24
    # Items GV: we'll define lists in main and pass last rows via defined names after creation.
    # Use SUM(C6:C23) style with constants matching item counts.

    n_fijos = 18
    n_var = 15
    gf_last_item = 5 + n_fijos  # 23
    gv_last_item = 5 + n_var    # 20

    f_fijos = {}
    f_var = {}
    for m in range(1, 13):
        col = get_column_letter(2 + m)  # C=3 for month 1... wait
        # GF: month 1 is column C = 3, month m is column 2+m
        gf_col = get_column_letter(2 + m)
        f_fijos[m] = f"={GF}!{gf_col}6:{gf_col}{gf_last_item}"
        # that's a range not a sum
        f_fijos[m] = f"=SUM({GF}!{gf_col}6:{gf_col}{gf_last_item})"
        f_var[m] = f"=SUM({GV}!{gf_col}6:{gf_col}{gv_last_item})"

    write_sumifs_row(ER_GAS_FIJOS, "Gastos fijos", f_fijos)
    write_sumifs_row(ER_GAS_COMIS, "Comisiones de vendedores", f_comis)
    write_sumifs_row(ER_GAS_VAR, "Gastos variables (sin comisiones)", f_var)

    set_row(ER_GAS_TOTAL, "TOTAL GASTOS OPERACIONALES", "t", PatternFill("solid", fgColor="4A5568"), True)
    ws.cell(ER_GAS_TOTAL, 1).font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    for m in range(1, 13):
        c = month_col(m)
        col = get_column_letter(c)
        cell = ws.cell(ER_GAS_TOTAL, c)
        cell.value = f"={col}{ER_GAS_FIJOS}+{col}{ER_GAS_COMIS}+{col}{ER_GAS_VAR}"
        cell.fill = PatternFill("solid", fgColor="4A5568")
        cell.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
        money(cell)
    ws.cell(ER_GAS_TOTAL, 14).value = f"=SUM(B{ER_GAS_TOTAL}:M{ER_GAS_TOTAL})"
    ws.cell(ER_GAS_TOTAL, 14).fill = PatternFill("solid", fgColor="4A5568")
    ws.cell(ER_GAS_TOTAL, 14).font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    money(ws.cell(ER_GAS_TOTAL, 14))

    set_row(ER_UTIL_OP, "UTILIDAD OPERACIONAL", "t", fill_gold, True)
    for m in range(1, 13):
        c = month_col(m)
        col = get_column_letter(c)
        cell = ws.cell(ER_UTIL_OP, c)
        cell.value = f"={col}{ER_MARGEN}-{col}{ER_GAS_TOTAL}"
        cell.fill = fill_gold
        cell.font = Font(name="Calibri", size=11, bold=True, color=NAVY_DEEP)
        money(cell)
    ws.cell(ER_UTIL_OP, 14).value = f"=N{ER_MARGEN}-N{ER_GAS_TOTAL}"
    ws.cell(ER_UTIL_OP, 14).fill = fill_gold
    ws.cell(ER_UTIL_OP, 1).fill = fill_gold
    money(ws.cell(ER_UTIL_OP, 14))

    set_row(ER_UTIL_PCT, "Margen operacional %", "p")
    for m in range(1, 14):
        src_c = m + 1 if m < 13 else 14
        col = get_column_letter(src_c)
        cell = ws.cell(ER_UTIL_PCT, src_c)
        cell.value = f'=IF({col}{ER_ING_TOTAL}=0,"",{col}{ER_UTIL_OP}/{col}{ER_ING_TOTAL})'
        pct(cell)
        cell.fill = fill_formula
        cell.border = thin

    section_row(ws, 30, 14, "IMPUESTO A LA RENTA ESTIMADO (no es el F29)")
    set_row(ER_IMP_RENTA, "Provisión 1ª categoría del mes", "f")
    for m in range(1, 13):
        c = month_col(m)
        col = get_column_letter(c)
        cell = ws.cell(ER_IMP_RENTA, c)
        cell.value = f"=IF({col}{ER_UTIL_OP}>0,ROUND({col}{ER_UTIL_OP}*{RENTA},0),0)"
        cell.fill = fill_formula
        money(cell)
    ws.cell(ER_IMP_RENTA, 14).value = f"=SUM(B{ER_IMP_RENTA}:M{ER_IMP_RENTA})"
    money(ws.cell(ER_IMP_RENTA, 14))
    ws.cell(ER_IMP_RENTA, 14).fill = fill_gold

    set_row(ER_UTIL_NETA, "UTILIDAD NETA ESTIMADA", "t", fill_deep, True)
    ws.cell(ER_UTIL_NETA, 1).font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    for m in range(1, 13):
        c = month_col(m)
        col = get_column_letter(c)
        cell = ws.cell(ER_UTIL_NETA, c)
        cell.value = f"={col}{ER_UTIL_OP}-{col}{ER_IMP_RENTA}"
        cell.fill = fill_deep
        cell.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
        money(cell)
    ws.cell(ER_UTIL_NETA, 14).value = f"=N{ER_UTIL_OP}-N{ER_IMP_RENTA}"
    ws.cell(ER_UTIL_NETA, 14).fill = fill_deep
    ws.cell(ER_UTIL_NETA, 14).font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    money(ws.cell(ER_UTIL_NETA, 14))

    ws.column_dimensions["A"].width = 38
    for i in range(2, 15):
        ws.column_dimensions[get_column_letter(i)].width = 13
    ws.freeze_panes = "B5"
    ws.sheet_properties.tabColor = "2B6CB0"

    ws.conditional_formatting.add(
        f"B{ER_UTIL_OP}:N{ER_UTIL_OP}",
        CellIsRule(operator="lessThan", formula=["0"], fill=fill_red, font=Font(color=RED, bold=True)),
    )


def build_iva(ws):
    page_setup(ws)
    header_bar(ws, 1, 14, "  IVA mensual")
    sub_bar(ws, 2, 14, "  Débito de ventas − crédito de compras. El remanente se arrastra. No edites las fórmulas.")
    ws.merge_cells("A3:N3")
    ws.cell(3, 1, "El IVA se declara en el F29 del mes siguiente (normalmente hasta el día 12).").font = font_muted
    fill_month_headers(ws, 4)

    rows = {
        6: "IVA DÉBITO",
        7: "Débito ventas de autos / comisiones",
        8: "Débito taller y repuestos",
        9: "Débito otros ingresos",
        10: "TOTAL DÉBITO",
        12: "IVA CRÉDITO",
        13: "Crédito compra de autos",
        14: "Crédito costos de taller/repuestos",
        15: "Crédito gastos fijos",
        16: "Crédito gastos variables",
        17: "TOTAL CRÉDITO",
        19: "IVA determinado (débito − crédito)",
        20: "Remanente de crédito del mes anterior",
        21: "IVA A PAGAR",
        22: "Remanente para el mes siguiente",
    }

    for r, lab in rows.items():
        ws.cell(r, 1, lab).alignment = left
        ws.cell(r, 1).border = thin
        ws.cell(r, 1).font = font_label if r in (6, 12, 21) else font_normal

    section_row(ws, 6, 14, "IVA DÉBITO")
    section_row(ws, 12, 14, "IVA CRÉDITO")

    def put(row, month_formulas, fill=fill_formula, total=True):
        for m in range(1, 13):
            cell = ws.cell(row, month_col(m), month_formulas[m])
            cell.fill = fill
            cell.font = font_formula
            cell.border = thin
            money(cell)
        if total:
            cell = ws.cell(row, 14, f"=SUM(B{row}:M{row})")
            cell.fill = fill_gold
            cell.border = thin
            money(cell)

    deb_auto, deb_tr, deb_oi = {}, {}, {}
    cre_auto, cre_tr, cre_gf, cre_gv = {}, {}, {}, {}
    for m in range(1, 13):
        deb_auto[m] = (
            f'=SUMIFS({VA}!$X${VA_FIRST}:$X${VA_LAST},{VA}!$U${VA_FIRST}:$U${VA_LAST},{ANIO},'
            f'{VA}!$V${VA_FIRST}:$V${VA_LAST},{m},{VA}!$Q${VA_FIRST}:$Q${VA_LAST},"Vendido")'
        )
        deb_tr[m] = (
            f'=SUMIFS({TR}!$G${TR_FIRST}:$G${TR_LAST},{TR}!$M${TR_FIRST}:$M${TR_LAST},{ANIO},'
            f'{TR}!$N${TR_FIRST}:$N${TR_LAST},{m})'
        )
        deb_oi[m] = (
            f'=SUMIFS({OI}!$F${OI_FIRST}:$F${OI_LAST},{OI}!$H${OI_FIRST}:$H${OI_LAST},{ANIO},'
            f'{OI}!$I${OI_FIRST}:$I${OI_LAST},{m})'
        )
        cre_auto[m] = (
            f'=SUMIFS({VA}!$L${VA_FIRST}:$L${VA_LAST},{VA}!$S${VA_FIRST}:$S${VA_LAST},{ANIO},'
            f'{VA}!$T${VA_FIRST}:$T${VA_LAST},{m},{VA}!$J${VA_FIRST}:$J${VA_LAST},"Sí")'
        )
        cre_tr[m] = (
            f'=SUMIFS({TR}!$K${TR_FIRST}:$K${TR_LAST},{TR}!$M${TR_FIRST}:$M${TR_LAST},{ANIO},'
            f'{TR}!$N${TR_FIRST}:$N${TR_LAST},{m})'
        )
        cre_gf[m] = suma_iva_tipo("Gasto fijo", m)
        cre_gv[m] = suma_iva_tipo("Gasto variable", m)

    ws.cell(7, 1, "Débito ventas de autos / comisiones")
    put(7, deb_auto)
    ws.cell(8, 1, "Débito taller y repuestos")
    put(8, deb_tr)
    ws.cell(9, 1, "Débito otros ingresos")
    put(9, deb_oi)

    ws.cell(10, 1, "TOTAL DÉBITO").font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    for m in range(1, 14):
        c = m + 1 if m < 13 else 14
        col = get_column_letter(c)
        if m < 13:
            ws.cell(10, c).value = f"={col}7+{col}8+{col}9"
        else:
            ws.cell(10, c).value = "=SUM(B10:M10)"
        ws.cell(10, c).fill = fill_navy
        ws.cell(10, c).font = Font(name="Calibri", size=10, bold=True, color=WHITE)
        ws.cell(10, c).border = thin
        money(ws.cell(10, c))
    ws.cell(10, 1).fill = fill_navy
    ws.cell(10, 1).font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    ws.cell(10, 1).border = thin

    ws.cell(13, 1, "Crédito compra de autos")
    put(13, cre_auto)
    ws.cell(14, 1, "Crédito costos de taller/repuestos")
    put(14, cre_tr)
    ws.cell(15, 1, "Crédito gastos fijos")
    put(15, cre_gf)
    ws.cell(16, 1, "Crédito gastos variables")
    put(16, cre_gv)

    ws.cell(17, 1, "TOTAL CRÉDITO").font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    for m in range(1, 14):
        c = m + 1 if m < 13 else 14
        col = get_column_letter(c)
        if m < 13:
            ws.cell(17, c).value = f"={col}13+{col}14+{col}15+{col}16"
        else:
            ws.cell(17, c).value = "=SUM(B17:M17)"
        ws.cell(17, c).fill = PatternFill("solid", fgColor="4A5568")
        ws.cell(17, c).font = Font(name="Calibri", size=10, bold=True, color=WHITE)
        ws.cell(17, c).border = thin
        money(ws.cell(17, c))
    ws.cell(17, 1).fill = PatternFill("solid", fgColor="4A5568")
    ws.cell(17, 1).font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    ws.cell(17, 1).border = thin

    ws.cell(19, 1, "IVA determinado (débito − crédito)")
    for m in range(1, 13):
        col = get_column_letter(month_col(m))
        formula_cell(ws.cell(19, month_col(m)), f"={col}10-{col}17", "clp")
    formula_cell(ws.cell(19, 14), "=SUM(B19:M19)", "clp")
    ws.cell(19, 14).fill = fill_gold

    ws.cell(20, 1, "Remanente de crédito del mes anterior")
    formula_cell(ws.cell(20, 2), "=0", "clp")  # enero: se puede sobreescribir con input
    ws.cell(20, 2).fill = fill_input  # allow opening remanente
    ws.cell(20, 2).font = font_input
    for m in range(2, 13):
        prev = get_column_letter(month_col(m - 1))
        formula_cell(ws.cell(20, month_col(m)), f"={prev}22", "clp")
    formula_cell(ws.cell(20, 14), "=-", )  # unused
    ws.cell(20, 14).value = ""
    ws.cell(20, 14).fill = fill_white
    ws.cell(20, 1).comment = Comment(
        "En enero puedes escribir el remanente que venía de diciembre del año anterior.",
        "Sistema",
    )

    ws.cell(21, 1, "IVA A PAGAR").font = Font(name="Calibri", size=11, bold=True, color=WHITE)
    ws.cell(21, 1).fill = fill_navy
    for m in range(1, 13):
        col = get_column_letter(month_col(m))
        cell = ws.cell(21, month_col(m))
        cell.value = f"=MAX(0,{col}19-{col}20)"
        cell.fill = fill_navy
        cell.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
        cell.border = thin
        money(cell)
    ws.cell(21, 14).value = "=SUM(B21:M21)"
    ws.cell(21, 14).fill = fill_gold
    ws.cell(21, 14).font = Font(name="Calibri", size=10, bold=True, color=NAVY_DEEP)
    money(ws.cell(21, 14))
    ws.cell(21, 14).border = thin

    ws.cell(22, 1, "Remanente para el mes siguiente")
    for m in range(1, 13):
        col = get_column_letter(month_col(m))
        formula_cell(ws.cell(22, month_col(m)), f"=MAX(0,{col}20-{col}19)", "clp")
    ws.cell(22, 14).value = ""
    ws.cell(22, 14).border = thin

    ws.column_dimensions["A"].width = 42
    for i in range(2, 15):
        ws.column_dimensions[get_column_letter(i)].width = 13
    ws.freeze_panes = "B5"
    ws.sheet_properties.tabColor = "C53030"

    ws.cell(24, 1, "Cómo leerlo").font = font_label
    ws.merge_cells("A25:N26")
    ws.cell(
        25,
        1,
        "Si IVA a pagar > 0, ese monto va en el F29 junto con el PPM. Si remanente > 0, no pagas IVA ese mes y el saldo baja el IVA de los meses siguientes. "
        "El crédito de autos entra en el mes de la COMPRA (fecha compra), el débito en el mes de la VENTA.",
    ).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[25].height = 40


def build_ppm(ws):
    page_setup(ws)
    header_bar(ws, 1, 14, "  PPM e impuesto a la renta")
    sub_bar(ws, 2, 14, "  PPM = ingresos netos del mes × tasa. Se paga en el F29. No es un gasto; es un anticipo del impuesto anual.")
    fill_month_headers(ws, 4)

    labels = {
        6: "Base PPM (ingresos netos del mes)",
        7: "Tasa PPM",
        8: "PPM DEL MES",
        9: "PPM acumulado del año",
        11: "Utilidad operacional del mes",
        12: "Provisión 1ª categoría del mes",
        13: "PPM del mes (anticipo)",
        14: "Saldo renta del mes (provisión − PPM)",
        16: "Provisión 1ª categoría acumulada",
        17: "PPM acumulado",
        18: "SALDO ESTIMADO RENTA DEL AÑO",
    }
    for r, lab in labels.items():
        ws.cell(r, 1, lab).font = font_normal
        ws.cell(r, 1).alignment = left
        ws.cell(r, 1).border = thin

    section_row(ws, 5, 14, "PAGO PROVISIONAL MENSUAL")
    for m in range(1, 13):
        col = get_column_letter(month_col(m))
        formula_cell(ws.cell(6, month_col(m)), f"={ER}!{col}{ER_ING_TOTAL}", "clp")
        formula_cell(ws.cell(7, month_col(m)), f"={PPM}", "pct")
        cell = ws.cell(8, month_col(m))
        cell.value = f"=ROUND({col}6*{col}7,0)"
        cell.fill = fill_navy
        cell.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
        cell.border = thin
        money(cell)
        if m == 1:
            formula_cell(ws.cell(9, 2), "=B8", "clp")
        else:
            prev = get_column_letter(month_col(m - 1))
            formula_cell(ws.cell(9, month_col(m)), f"={prev}9+{col}8", "clp")

        formula_cell(ws.cell(11, month_col(m)), f"={ER}!{col}{ER_UTIL_OP}", "clp")
        formula_cell(ws.cell(12, month_col(m)), f"={ER}!{col}{ER_IMP_RENTA}", "clp")
        formula_cell(ws.cell(13, month_col(m)), f"={col}8", "clp")
        formula_cell(ws.cell(14, month_col(m)), f"={col}12-{col}13", "clp")

    for r in (6, 8, 9, 11, 12, 13, 14, 16, 17, 18):
        cell = ws.cell(r, 14, f"=SUM(B{r}:M{r})" if r not in (9, 16, 17, 18) else f"=M{r}")
        if r == 6:
            cell.value = f"=N{ER_ING_TOTAL}" if False else "=SUM(B6:M6)"
        if r == 8:
            cell.value = "=SUM(B8:M8)"
        if r == 9:
            cell.value = "=M9"
        if r == 11:
            cell.value = "=SUM(B11:M11)"
        if r == 12:
            cell.value = "=SUM(B12:M12)"
        if r == 13:
            cell.value = "=SUM(B13:M13)"
        if r == 14:
            cell.value = "=N12-N13"
        cell.fill = fill_gold
        cell.border = thin
        money(cell)

    ws.cell(7, 14).value = f"={PPM}"
    pct(ws.cell(7, 14))
    ws.cell(7, 14).fill = fill_gold
    ws.cell(7, 14).border = thin

    section_row(ws, 10, 14, "CRUCE CON IMPUESTO A LA RENTA (estimado)")
    section_row(ws, 15, 14, "CIERRE DE AÑO (acumulado)")
    formula_cell(ws.cell(16, 2), "=B12", "clp")
    formula_cell(ws.cell(17, 2), "=B8", "clp")
    for m in range(2, 13):
        col = get_column_letter(month_col(m))
        prev = get_column_letter(month_col(m - 1))
        formula_cell(ws.cell(16, month_col(m)), f"={prev}16+{col}12", "clp")
        formula_cell(ws.cell(17, month_col(m)), f"={prev}17+{col}8", "clp")
    for m in range(1, 13):
        col = get_column_letter(month_col(m))
        cell = ws.cell(18, month_col(m))
        cell.value = f"={col}16-{col}17"
        cell.fill = fill_deep
        cell.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
        cell.border = thin
        money(cell)
    ws.cell(18, 14).value = "=M18"
    ws.cell(18, 14).fill = fill_gold
    money(ws.cell(18, 14))
    ws.cell(16, 14).value = "=M16"
    ws.cell(17, 14).value = "=M17"
    money(ws.cell(16, 14))
    money(ws.cell(17, 14))
    ws.cell(16, 14).fill = fill_gold
    ws.cell(17, 14).fill = fill_gold
    ws.cell(16, 14).border = thin
    ws.cell(17, 14).border = thin

    ws.cell(18, 1).fill = fill_deep
    ws.cell(18, 1).font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    ws.cell(8, 1).fill = fill_navy
    ws.cell(8, 1).font = Font(name="Calibri", size=10, bold=True, color=WHITE)

    ws.merge_cells("A20:N22")
    ws.cell(
        20,
        1,
        "Si el SALDO ESTIMADO RENTA DEL AÑO es positivo, es lo que aproximadamente faltaría por pagar en abril (Renta). "
        "Si es negativo, el PPM se te habría ido de más y queda como saldo a favor. "
        "Es una estimación de gestión: el F22/F29 oficial lo hace el contador con el Formulario 22 al cierre. "
        "Recuerda: el PPM se calcula sobre ingresos, aunque el mes esté a pérdida.",
    ).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[20].height = 56

    ws.column_dimensions["A"].width = 42
    for i in range(2, 15):
        ws.column_dimensions[get_column_letter(i)].width = 13
    ws.freeze_panes = "B5"
    ws.sheet_properties.tabColor = "C53030"


def kpi_box(ws, r, c, title, formula, sub_formula, cols=2):
    ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + cols - 1)
    ws.merge_cells(start_row=r + 1, start_column=c, end_row=r + 1, end_column=c + cols - 1)
    ws.merge_cells(start_row=r + 2, start_column=c, end_row=r + 2, end_column=c + cols - 1)
    t = ws.cell(r, c, title)
    t.font = Font(name="Calibri", size=9, bold=True, color=WHITE)
    t.fill = fill_navy
    t.alignment = Alignment(horizontal="center", vertical="center")
    for i in range(cols):
        ws.cell(r, c + i).fill = fill_navy
        ws.cell(r, c + i).border = thick_gold
        ws.cell(r + 1, c + i).border = thick_gold
        ws.cell(r + 2, c + i).border = thick_gold
        ws.cell(r + 1, c + i).fill = fill_cream
        ws.cell(r + 2, c + i).fill = fill_cream
    val = ws.cell(r + 1, c, formula)
    val.font = font_kpi
    val.alignment = center
    money(val)
    sub = ws.cell(r + 2, c, sub_formula)
    sub.font = font_muted
    sub.alignment = center
    ws.row_dimensions[r].height = 18
    ws.row_dimensions[r + 1].height = 28
    ws.row_dimensions[r + 2].height = 18


def build_resumen(ws):
    page_setup(ws)
    header_bar(ws, 1, 12, "  Resumen de gestión")
    # period line
    ws.merge_cells("A2:L2")
    ws.cell(2, 1).value = f'= {NOMBRE} & "  ·  " & Configuracion!B16 & " " & {ANIO} & "  ·  Cifras netas, sin IVA"'
    ws.cell(2, 1).font = Font(name="Calibri", size=11, color=WHITE)
    for c in range(1, 13):
        ws.cell(2, c).fill = fill_deep
        ws.cell(2, c).font = Font(name="Calibri", size=11, color="E2E8F0")
    ws.row_dimensions[2].height = 22

    idx = f"INDEX({ER}!B{{row}}:{ER}!M{{row}},{MES})"
    def er(row):
        return f"INDEX({ER}!$B${row}:$M${row},1,{MES})"

    def iva(row):
        return f"INDEX({SH_IVA}!$B${row}:$M${row},1,{MES})"

    def ppm(row):
        return f"INDEX({SH_PPM}!$B${row}:$M${row},1,{MES})"

    kpi_box(ws, 4, 1, "INGRESOS DEL MES", f"={er(ER_ING_TOTAL)}", '="Mes en curso"', 3)
    kpi_box(ws, 4, 4, "MARGEN BRUTO", f"={er(ER_MARGEN)}", f'=IF({er(ER_ING_TOTAL)}=0,"",TEXT({er(ER_MARGEN_PCT)},"0.0%")&" sobre ventas")', 3)
    kpi_box(ws, 4, 7, "UTILIDAD OPERACIONAL", f"={er(ER_UTIL_OP)}", f'=IF({er(ER_UTIL_OP)}>=0,"La operación gana","La operación pierde")', 3)
    kpi_box(ws, 4, 10, "UTILIDAD NETA EST.", f"={er(ER_UTIL_NETA)}", '="Después de provisión 1ª cat."', 3)

    # Detalle ingresos / gastos
    ws.merge_cells("A8:F8")
    ws.cell(8, 1, "De dónde sale la plata este mes").font = Font(name="Calibri", size=12, bold=True, color=WHITE)
    for c in range(1, 7):
        ws.cell(8, c).fill = fill_section
        ws.cell(8, c).font = Font(name="Calibri", size=12, bold=True, color=WHITE)
    headers = ["Línea", "Monto neto", "% de ingresos"]
    for i, h in enumerate(headers, 1):
        cell = ws.cell(9, i, h)
        cell.font = font_h
        cell.fill = fill_navy
        cell.alignment = center
        cell.border = thin

    lineas = [
        (10, "Autos nuevos", ER_ING_NUEVOS),
        (11, "Autos usados", ER_ING_USADOS),
        (12, "Consignación (comisiones)", ER_ING_CONS),
        (13, "Taller", ER_ING_TALLER),
        (14, "Repuestos", ER_ING_REP),
        (15, "Otros ingresos", ER_ING_OTROS),
        (16, "TOTAL INGRESOS", ER_ING_TOTAL),
        (18, "Costo de autos vendidos", ER_COS_AUTOS),
        (19, "Costo taller y repuestos", ER_COS_TALLER),
        (20, "MARGEN BRUTO", ER_MARGEN),
        (22, "Gastos fijos", ER_GAS_FIJOS),
        (23, "Comisiones de vendedores", ER_GAS_COMIS),
        (24, "Gastos variables", ER_GAS_VAR),
        (25, "TOTAL GASTOS", ER_GAS_TOTAL),
        (27, "UTILIDAD OPERACIONAL", ER_UTIL_OP),
        (28, "Provisión impuesto renta", ER_IMP_RENTA),
        (29, "UTILIDAD NETA ESTIMADA", ER_UTIL_NETA),
    ]
    for r, lab, erow in lineas:
        ws.cell(r, 1, lab).border = thin
        ws.cell(r, 1).font = font_label if lab.isupper() else font_normal
        formula_cell(ws.cell(r, 2), f"={er(erow)}", "clp")
        formula_cell(ws.cell(r, 3), f'=IF({er(ER_ING_TOTAL)}=0,"",B{r}/{er(ER_ING_TOTAL)})', "pct")
        if lab.isupper():
            for c in range(1, 4):
                ws.cell(r, c).fill = fill_navy if "UTILIDAD NETA" not in lab else fill_deep
                ws.cell(r, c).font = Font(name="Calibri", size=10, bold=True, color=WHITE)
            if "MARGEN BRUTO" in lab:
                for c in range(1, 4):
                    ws.cell(r, c).fill = fill_green
                    ws.cell(r, c).font = Font(name="Calibri", size=10, bold=True, color=GREEN)
            if lab == "UTILIDAD OPERACIONAL":
                for c in range(1, 4):
                    ws.cell(r, c).fill = fill_gold
                    ws.cell(r, c).font = Font(name="Calibri", size=10, bold=True, color=NAVY_DEEP)

    # Tax box
    ws.merge_cells("H8:L8")
    ws.cell(8, 8, "Qué hay que pagar al SII este mes (F29)").font = Font(name="Calibri", size=12, bold=True, color=WHITE)
    for c in range(8, 13):
        ws.cell(8, c).fill = PatternFill("solid", fgColor="9B2C2C")
        ws.cell(8, c).font = Font(name="Calibri", size=12, bold=True, color=WHITE)

    tax_rows = [
        (9, "Campo", "Monto"),
        (10, "IVA débito", f"={iva(10)}"),
        (11, "IVA crédito", f"={iva(17)}"),
        (12, "IVA determinado", f"={iva(19)}"),
        (13, "Remanente anterior", f"={iva(20)}"),
        (14, "IVA A PAGAR", f"={iva(21)}"),
        (15, "PPM A PAGAR", f"={ppm(8)}"),
        (16, "TOTAL F29 (IVA + PPM)", f"=I14+I15"),
        (18, "Vencimiento (orientativo)", None),
    ]
    ws.cell(9, 8, "Concepto").font = font_h
    ws.cell(9, 8).fill = fill_navy
    ws.cell(9, 9, "Monto").font = font_h
    ws.cell(9, 9).fill = fill_navy
    ws.cell(9, 8).border = thin
    ws.cell(9, 9).border = thin
    ws.merge_cells("I9:L9")
    for c in range(9, 13):
        ws.cell(9, c).fill = fill_navy
        ws.cell(9, c).border = thin

    items = [
        (10, "IVA débito", f"={iva(10)}"),
        (11, "IVA crédito", f"={iva(17)}"),
        (12, "IVA determinado", f"={iva(19)}"),
        (13, "Remanente anterior", f"={iva(20)}"),
        (14, "IVA A PAGAR", f"={iva(21)}"),
        (15, "PPM A PAGAR", f"={ppm(8)}"),
        (16, "TOTAL F29 (IVA + PPM)", "=I14+I15"),
    ]
    for r, lab, fml in items:
        ws.merge_cells(start_row=r, start_column=9, end_row=r, end_column=12)
        ws.cell(r, 8, lab).border = thin
        ws.cell(r, 8).font = font_label if lab.isupper() else font_normal
        cell = ws.cell(r, 9, fml)
        formula_cell(cell, fml, "clp")
        for c in range(9, 13):
            ws.cell(r, c).border = thin
            ws.cell(r, c).fill = fill_formula
        if lab.startswith("TOTAL") or lab.startswith("IVA A PAGAR") or lab.startswith("PPM"):
            ws.cell(r, 8).fill = fill_navy
            ws.cell(r, 8).font = Font(name="Calibri", size=10, bold=True, color=WHITE)
            for c in range(9, 13):
                ws.cell(r, c).fill = fill_navy
                ws.cell(r, c).font = Font(name="Calibri", size=10, bold=True, color=WHITE)

    ws.merge_cells("H18:L19")
    ws.cell(18, 8).value = (
        f'="Declarar en el F29 de "&TEXT(DATE({ANIO},{MES},1)+32-DAY(DATE({ANIO},{MES},1)+32),"MMMM YYYY")'
        f'&". Fecha habitual: día 12 del mes siguiente (o hábil siguiente)."'
    )
    ws.cell(18, 8).alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(18, 8).font = font_muted
    ws.row_dimensions[18].height = 18
    ws.row_dimensions[19].height = 18

    # Acumulado año
    ws.merge_cells("A31:C31")
    ws.cell(31, 1, "Acumulado del año").font = Font(name="Calibri", size=12, bold=True, color=WHITE)
    for c in range(1, 4):
        ws.cell(31, c).fill = fill_section
        ws.cell(31, c).font = Font(name="Calibri", size=12, bold=True, color=WHITE)

    headers2 = ["Indicador", "Enero a la fecha", "Año completo"]
    # We'll show year total from column N and YTD as SUM of months 1..current
    ws.cell(32, 1, "Indicador").font = font_h
    ws.cell(32, 2, "Acumulado hasta el mes en curso").font = font_h
    ws.cell(32, 3, "Proyección año (suma 12 meses)").font = font_h
    for c in range(1, 4):
        ws.cell(32, c).fill = fill_navy
        ws.cell(32, c).font = font_h
        ws.cell(32, c).alignment = center
        ws.cell(32, c).border = thin

    def ytd(row):
        return f"=SUM(OFFSET({ER}!$B${row},0,0,1,{MES}))"

    ytd_rows = [
        (33, "Ingresos", ER_ING_TOTAL),
        (34, "Margen bruto", ER_MARGEN),
        (35, "Utilidad operacional", ER_UTIL_OP),
        (36, "Utilidad neta estimada", ER_UTIL_NETA),
        (37, "IVA pagado / a pagar", None),
        (38, "PPM pagado / a pagar", None),
    ]
    for r, lab, erow in ytd_rows:
        ws.cell(r, 1, lab).border = thin
        ws.cell(r, 1).font = font_normal
        if erow:
            formula_cell(ws.cell(r, 2), ytd(erow), "clp")
            formula_cell(ws.cell(r, 3), f"={ER}!N{erow}", "clp")
        elif lab.startswith("IVA"):
            formula_cell(ws.cell(r, 2), f"=SUM(OFFSET({SH_IVA}!$B$21,0,0,1,{MES}))", "clp")
            formula_cell(ws.cell(r, 3), f"={SH_IVA}!N21", "clp")
        else:
            formula_cell(ws.cell(r, 2), f"=SUM(OFFSET({SH_PPM}!$B$8,0,0,1,{MES}))", "clp")
            formula_cell(ws.cell(r, 3), f"={SH_PPM}!N8", "clp")

    # Punto de equilibrio
    ws.merge_cells("H31:L31")
    ws.cell(31, 8, "Lectura rápida").font = Font(name="Calibri", size=12, bold=True, color=WHITE)
    for c in range(8, 13):
        ws.cell(31, c).fill = fill_section
    ws.merge_cells("H32:L38")
    ws.cell(32, 8).value = (
        f'=IF({er(ER_ING_TOTAL)}=0,"Carga ventas o elige un mes con movimiento.",'
        f'IF({er(ER_UTIL_OP)}>=0,'
        f'"Este mes la automotora gana "&TEXT({er(ER_UTIL_OP)},"$#,##0")&" a nivel operacional. "'
        f'&"El F29 estimado (IVA + PPM) es "&TEXT(I16,"$#,##0")&".",'
        f'"Este mes la operación pierde "&TEXT(-({er(ER_UTIL_OP)}),"$#,##0")&". "'
        f'&"Revisa gastos fijos y el margen por auto. El F29 igual puede corresponder: "&TEXT(I16,"$#,##0")&"."))'
    )
    ws.cell(32, 8).alignment = Alignment(wrap_text=True, vertical="top", indent=1)
    ws.cell(32, 8).font = Font(name="Calibri", size=11, color=NAVY)
    ws.cell(32, 8).fill = fill_cream

    # Chart data is on ER — chart
    chart = BarChart()
    chart.type = "col"
    chart.grouping = "clustered"
    chart.title = "Ingresos, gastos y utilidad operacional"
    chart.y_axis.title = None
    chart.x_axis.title = None
    chart.style = 10
    chart.y_axis.numFmt = '"$"#,##0'
    chart.height = 8
    chart.width = 18

    cats = Reference(ws.parent[ER], min_col=2, min_row=ER_HDR, max_col=13, max_row=ER_HDR)
    # We need to add series from ER sheet
    data_ing = Reference(ws.parent[ER], min_col=1, min_row=ER_ING_TOTAL, max_col=13, max_row=ER_ING_TOTAL)
    data_gas = Reference(ws.parent[ER], min_col=1, min_row=ER_GAS_TOTAL, max_col=13, max_row=ER_GAS_TOTAL)
    data_uti = Reference(ws.parent[ER], min_col=1, min_row=ER_UTIL_OP, max_col=13, max_row=ER_UTIL_OP)
    chart.add_data(data_ing, from_rows=True, titles_from_data=True)
    chart.add_data(data_gas, from_rows=True, titles_from_data=True)
    chart.add_data(data_uti, from_rows=True, titles_from_data=True)
    chart.set_categories(cats)
    chart.shape = 4
    chart.legend.position = "b"
    ws.add_chart(chart, "A40")

    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 14
    for col in "DEFG":
        ws.column_dimensions[col].width = 14
    ws.column_dimensions["H"].width = 26
    for col in "IJKL":
        ws.column_dimensions[col].width = 13

    ws.freeze_panes = "A4"
    ws.sheet_properties.tabColor = NAVY
    ws.row_dimensions[1].height = 30

    ws.conditional_formatting.add(
        "B27",
        CellIsRule(operator="lessThan", formula=["0"], fill=fill_red, font=Font(color=RED, bold=True)),
    )


def main():
    wb = Workbook()

    ws_comp = wb.active
    ws_comp.title = "Comprobantes"
    ws_in = wb.create_sheet("Instrucciones")
    ws_cfg = wb.create_sheet("Configuracion")
    ws_res = wb.create_sheet("Resumen")
    ws_er = wb.create_sheet(ER)
    ws_va = wb.create_sheet(VA)
    ws_tr = wb.create_sheet(TR)
    ws_oi = wb.create_sheet(OI)
    ws_gf = wb.create_sheet(GF)
    ws_gv = wb.create_sheet(GV)
    ws_iva = wb.create_sheet(SH_IVA)
    ws_ppm = wb.create_sheet(SH_PPM)

    fijos = [(nombre, AFECTO_POR_DEFECTO[nombre]) for nombre in GASTOS_FIJOS]
    fijos_agosto = {
        "Arriendo local / patio": 2800000,
        "Sueldos (líquidos)": 4200000,
        "Leyes sociales empleador": 1150000,
        "Honorarios contador": 250000,
        "Seguros (local, stock, RC)": 180000,
        "Electricidad": 220000,
        "Internet y telefonía": 89000,
        "Software / facturación electrónica": 65000,
        "Publicidad fija (portales, redes)": 350000,
        "Alarma / seguridad": 45000,
        "Aseo": 80000,
    }

    variables = [(nombre, AFECTO_POR_DEFECTO[nombre]) for nombre in GASTOS_VARIABLES]
    var_agosto = {
        "Publicidad de campañas": 280000,
        "Preparación y detailing": 190000,
        "Combustible": 85000,
        "Comisiones Transbank / Webpay": 45000,
        "Traslado / flete de vehículos": 70000,
    }

    preparar_hoja_comprobantes(ws_comp)
    for concepto, monto in {**fijos_agosto, **var_agosto}.items():
        tipo = "Gasto fijo" if concepto in fijos_agosto else "Gasto variable"
        _escribir_comprobante(
            ws_comp,
            Movimiento(
                fecha=date(2026, 8, 15),
                tipo=tipo,
                categoria=concepto,
                descripcion="EJEMPLO — se puede borrar",
                neto=monto,
                afecto_iva=AFECTO_POR_DEFECTO[concepto],
                origen="Ejemplo",
                estado="Confirmado",
            ),
        )
    build_instrucciones(ws_in)
    build_config(ws_cfg)
    build_ventas(ws_va)
    build_taller(ws_tr)
    build_otros(ws_oi)
    build_gastos(
        ws_gf,
        "Gastos fijos",
        "Costos que se pagan aunque no vendas. El monto del mes es la suma de Comprobantes.",
        fijos,
        tab="DD6B20",
    )
    build_gastos(
        ws_gv,
        "Gastos variables",
        "Costos que suben cuando hay más movimiento. Las comisiones de vendedores se toman de Ventas_Autos, no las dupliques aquí.",
        variables,
        tab="DD6B20",
    )
    build_er(ws_er)
    build_iva(ws_iva)
    build_ppm(ws_ppm)
    build_resumen(ws_res)

    # Defined names
    for name, ref in {
        "TasaIVA": "Configuracion!$B$11",
        "TasaPPM": "Configuracion!$B$12",
        "TasaRenta": "Configuracion!$B$13",
        "AnioTrabajo": "Configuracion!$B$7",
        "MesTrabajo": "Configuracion!$B$8",
    }.items():
        wb.defined_names.add(DefinedName(name=name, attr_text=ref))

    # print titles
    for ws in wb.worksheets:
        ws.oddHeader.left.text = "Control financiero automotora"
        ws.oddFooter.right.text = "Página &P de &N"

    wb.properties.title = "Control Financiero Automotora"
    wb.properties.creator = "Cursor"
    wb.properties.description = "Ingresos, gastos fijos y variables, IVA y PPM para automotora (Chile)."

    wb.save(OUT)
    print(f"OK {OUT}")


if __name__ == "__main__":
    main()
